"""Preliminary Philippine Building BOM & quantity-takeoff desktop app.

Run:  py -3.11 philippine_preliminary_bom.py
Optional Excel export dependency: pip install openpyxl
This is an estimating aid only; see the persistent disclaimer in the UI.
"""
from __future__ import annotations

import json
import math
from collections import defaultdict
from dataclasses import dataclass, asdict, field
from datetime import date
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, simpledialog

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.chart import PieChart, Reference
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

APP_TITLE = "Preliminary Building BOM & Takeoff | Philippines"
VERSION = "1.0"
BAR_WEIGHT = {6: .222, 8: .395, 10: .617, 12: .888, 16: 1.58, 20: 2.47, 25: 3.85, 28: 4.83, 32: 6.31} # kg/m
DEFAULTS = {"concrete_waste": 5, "rebar_waste": 10, "chb_waste": 5, "finish_waste": 10,
            "formwork_waste": 5, "cement_per_m3": 9, "sand_per_m3": .50, "gravel_per_m3": 1.0,
            "block_per_m2": 12.5, "mortar_m3_per_m2": .012, "plaster_m3_per_m2": .015,
            "paint_m2_per_litre": 10, "stock_bar_length": 6.0}

SCHEMAS = {
 "Footings": [("ID","F-1"),("Qty",1),("Length m",1.5),("Width m",1.5),("Thick m",.35),("Excav depth m",.8),("Work space m",.15),("Cover mm",75),("Bot X dia mm",12),("Bot X spacing mm",150),("Bot Y dia mm",12),("Bot Y spacing mm",150),("Dowel qty",4),("Dowel dia mm",16),("Dowel embed m",.8)],
 "Columns": [("ID","C-1"),("Qty",4),("Width m",.3),("Depth m",.3),("Height/storey m",3.0),("Storeys",2),("Cover mm",40),("Main bar qty",4),("Main dia mm",16),("Tie dia mm",10),("Tie spacing mm",150),("Hook m",.2),("Lap allowance m",.6)],
 "Tie Beams": [("ID","TB-1"),("Qty",4),("Length m",4),("Width m",.25),("Depth m",.35),("Cover mm",40),("Top qty",2),("Top dia mm",16),("Bottom qty",2),("Bottom dia mm",16),("Stirrup dia mm",10),("Stirrup spacing mm",150)],
 "Elevated Beams": [("ID","B-1"),("Qty",2),("Length m",4),("Width m",.25),("Depth m",.4),("Cover mm",40),("Top qty",2),("Top dia mm",16),("Bottom qty",2),("Bottom dia mm",16),("Stirrup dia mm",10),("Stirrup spacing mm",150)],
 "Slabs": [("ID","SOG-1"),("Type","Slab-on-grade"),("Qty/Floors",1),("Length m",8),("Width m",6),("Thickness m",.1),("Bar X dia mm",10),("Bar X spacing mm",200),("Bar Y dia mm",10),("Bar Y spacing mm",200),("Gravel base m",.1),("Sand bedding m",.05),("Vapor barrier","Yes"),("Finish thick m",.01)],
 "Suspended Slabs": [("ID","SS-1"),("Qty/Floors",1),("Length m",6),("Width m",5),("Thickness m",.125),("Clear span m",5),("Cover mm",25),("Bottom X dia mm",12),("Bottom X spacing mm",200),("Bottom Y dia mm",12),("Bottom Y spacing mm",200),("Top X dia mm",10),("Top X spacing mm",200),("Top Y dia mm",10),("Top Y spacing mm",200),("Finish thick m",.01)],
 "Walls": [("ID","W-1"),("Type","CHB"),("Qty",1),("Length m",8),("Height m",3),("Thickness m",.1),("Opening area m2",3),("Both-side plaster","Yes")],
 "Openings": [("ID","D-1"),("Type","Door"),("Qty",1),("Width m",.9),("Height m",2.1),("Material","Wood")],
 "Painting": [("ID","PAINT-1"),("Painted sides",2),("Finish coats",2),("Paint coverage m2/L",10),("Primer coverage m2/L",10),("Waste %",10)],
 "Tiles": [("ID","TILE-1"),("Location","Floor"),("Area m2",48),("Tile length mm",600),("Tile width mm",600),("Waste %",10),("Adhesive kg/m2",4),("Grout kg/m2",.25)],
 "Finishes": [("ID","Floor finish"),("Type","Floor"),("Area m2",48),("Coverage m2/unit",1),("Waste %",10)]
}

def f(row, key, default=0.0):
    try: return float(row.get(key, default))
    except (ValueError, TypeError): return default
def n(row, key, default=0): return max(0, f(row,key,default))
def yes(row,key): return str(row.get(key, "")).lower() in ("yes","true","1")
def wt(dia, length): return BAR_WEIGHT.get(round(dia), (dia*dia/162)) * length
def with_waste(q, pct): return q * (1 + pct/100)

@dataclass
class Project:
    info: dict = field(default_factory=lambda: {"Project name":"Sample Two-storey Building", "Client":"", "Location":"Philippines", "Building type":"Residential", "Storeys":2, "Estimator":"", "Date":str(date.today()), "Concrete strength":"f'c = 21 MPa (reference)", "Steel grade":"Grade 415 MPa (reference)"})
    assumptions: dict = field(default_factory=lambda: DEFAULTS.copy())
    items: dict = field(default_factory=lambda: {k: [dict(v) for v in []] for k in SCHEMAS})

class Calculator:
    """Transparent preliminary takeoff equations; dimensions are metric."""
    def __init__(self, project): self.p=project; self.a=project.assumptions; self.lines=[]; self.rebar=defaultdict(lambda:[0.,0.])
    def add_rebar(self, dia, length, note):
        if dia <= 0 or length <= 0:return
        length=with_waste(length,self.a["rebar_waste"]); self.rebar[round(dia)][0]+=length; self.rebar[round(dia)][1]+=wt(dia,length)
        self.lines.append(("Reinforcement",note,"Rebar %d mm"%round(dia),length,"m"))
    def add(self, cat, desc, qty, unit):
        self.lines.append((cat, desc, desc, max(0,qty), unit))
    def calc(self):
        for r in self.p.items["Footings"]: self.footing(r)
        for r in self.p.items["Columns"]: self.column(r)
        for r in self.p.items["Tie Beams"]: self.beam(r,"Tie beam")
        for r in self.p.items["Elevated Beams"]: self.beam(r,"Elevated beam")
        for r in self.p.items["Slabs"]: self.slab(r)
        for r in self.p.items.get("Suspended Slabs",[]): self.suspended_slab(r)
        for r in self.p.items["Walls"]: self.wall(r)
        for r in self.p.items["Openings"]: self.opening(r)
        for r in self.p.items.get("Painting",[]): self.painting(r)
        for r in self.p.items.get("Tiles",[]): self.tiles(r)
        for r in self.p.items["Finishes"]: self.finish(r)
        for dia,(ln,kg) in self.rebar.items():
            self.lines.append(("Reinforcement","Aggregated bars","Rebar %d mm"%dia,kg,"kg"))
            self.lines.append(("Reinforcement","Stock bars","%d mm / %.1f m"%(dia,self.a["stock_bar_length"]),math.ceil(ln/self.a["stock_bar_length"]),"pcs"))
        return self.lines
    def concrete(self, volume, note):
        v=with_waste(volume,self.a["concrete_waste"]); self.add("Concrete",note,v,"m³")
        self.add("Concrete materials",note+" cement",v*self.a["cement_per_m3"],"bags")
        self.add("Concrete materials",note+" sand",v*self.a["sand_per_m3"],"m³")
        self.add("Concrete materials",note+" gravel",v*self.a["gravel_per_m3"],"m³")
    def footing(self,r):
        q=n(r,"Qty"); L=n(r,"Length m"); W=n(r,"Width m"); T=n(r,"Thick m"); ws=n(r,"Work space m")
        self.add("Earthworks",r["ID"]+" excavation",q*(L+2*ws)*(W+2*ws)*n(r,"Excav depth m"),"m³")
        self.concrete(q*L*W*T,r["ID"]+" footing")
        self.add("Formwork",r["ID"]+" sides",with_waste(q*2*(L+W)*T,self.a["formwork_waste"]),"m²")
        cov=n(r,"Cover mm")/1000
        for direction in ("X","Y"):
            span=W if direction=="X" else L; run=L if direction=="X" else W
            count=math.floor(max(0,span-2*cov)/(n(r,"Bot %s spacing mm"%direction)/1000))+1 if n(r,"Bot %s spacing mm"%direction) else 0
            self.add_rebar(n(r,"Bot %s dia mm"%direction), q*count*max(0,run-2*cov), r["ID"]+" bottom "+direction)
        self.add_rebar(n(r,"Dowel dia mm"),q*n(r,"Dowel qty")*n(r,"Dowel embed m"),r["ID"]+" dowels")
    def column(self,r):
        q=n(r,"Qty")*n(r,"Storeys"); w=n(r,"Width m"); d=n(r,"Depth m"); h=n(r,"Height/storey m"); cov=n(r,"Cover mm")/1000
        self.concrete(q*w*d*h,r["ID"]+" columns"); self.add("Formwork",r["ID"]+" columns",with_waste(q*2*(w+d)*h,self.a["formwork_waste"]),"m²")
        self.add_rebar(n(r,"Main dia mm"),q*n(r,"Main bar qty")*(h+n(r,"Lap allowance m")),r["ID"]+" vertical")
        cnt=math.ceil(h/(n(r,"Tie spacing mm")/1000))+1 if n(r,"Tie spacing mm") else 0
        self.add_rebar(n(r,"Tie dia mm"),q*cnt*(2*max(0,w-2*cov)+2*max(0,d-2*cov)+n(r,"Hook m")),r["ID"]+" ties")
    def beam(self,r,kind):
        q=n(r,"Qty"); L=n(r,"Length m"); w=n(r,"Width m"); d=n(r,"Depth m"); cov=n(r,"Cover mm")/1000
        self.concrete(q*L*w*d,r["ID"]+" "+kind); self.add("Formwork",r["ID"]+" "+kind,with_waste(q*L*(2*d+w),self.a["formwork_waste"]),"m²")
        self.add_rebar(n(r,"Top dia mm"),q*n(r,"Top qty")*L,r["ID"]+" top bars"); self.add_rebar(n(r,"Bottom dia mm"),q*n(r,"Bottom qty")*L,r["ID"]+" bottom bars")
        cnt=math.ceil(L/(n(r,"Stirrup spacing mm")/1000))+1 if n(r,"Stirrup spacing mm") else 0
        self.add_rebar(n(r,"Stirrup dia mm"),q*cnt*(2*max(0,w-2*cov)+2*max(0,d-2*cov)+.2),r["ID"]+" stirrups")
    def slab(self,r):
        q=n(r,"Qty/Floors"); L=n(r,"Length m"); W=n(r,"Width m"); A=q*L*W
        self.concrete(A*n(r,"Thickness m"),r["ID"]+" slab")
        for x in ("X","Y"):
            span=W if x=="X" else L; run=L if x=="X" else W; space=n(r,"Bar %s spacing mm"%x)/1000
            self.add_rebar(n(r,"Bar %s dia mm"%x),q*(math.floor(span/space)+1)*run if space else 0,r["ID"]+" slab "+x)
        self.add("Flooring",r["ID"]+" gravel base",A*n(r,"Gravel base m"),"m³"); self.add("Flooring",r["ID"]+" sand bedding",A*n(r,"Sand bedding m"),"m³")
        if yes(r,"Vapor barrier"): self.add("Flooring",r["ID"]+" vapor barrier",A,"m²")
        self.add("Finishes",r["ID"]+" floor finish",A*n(r,"Finish thick m"),"m³")
    def suspended_slab(self,r):
        """Elevated RC slab: concrete, underside/edge formwork and two-layer bar mat."""
        q=n(r,"Qty/Floors"); L=n(r,"Length m"); W=n(r,"Width m"); A=q*L*W
        self.concrete(A*n(r,"Thickness m"),r["ID"]+" suspended slab")
        # Soffit plus perimeter slab edges. Supporting beams are taken off separately.
        self.add("Formwork",r["ID"]+" soffit and slab edges",with_waste(A + q*2*(L+W)*n(r,"Thickness m"),self.a["formwork_waste"]),"m²")
        for layer in ("Bottom","Top"):
            for axis in ("X","Y"):
                span=W if axis=="X" else L; run=L if axis=="X" else W; space=n(r,"%s %s spacing mm"%(layer,axis))/1000
                count=math.floor(span/space)+1 if space else 0
                self.add_rebar(n(r,"%s %s dia mm"%(layer,axis)),q*count*run,r["ID"]+" %s %s bars"%(layer.lower(),axis))
        self.add("Flooring",r["ID"]+" suspended floor area",A,"m²")
        self.add("Finishes",r["ID"]+" floor finish",A*n(r,"Finish thick m"),"m³")
    def wall(self,r):
        q=n(r,"Qty"); gross=q*n(r,"Length m")*n(r,"Height m"); net=max(0,gross-n(r,"Opening area m2")); typ=str(r.get("Type",""))
        self.add("Masonry",r["ID"]+" gross wall area",gross,"m²"); self.add("Masonry",r["ID"]+" net wall area",net,"m²")
        self.add("Masonry",r["ID"]+" wall volume",net*n(r,"Thickness m"),"m³")
        if "chb" in typ.lower() or "block" in typ.lower(): self.add("Masonry",r["ID"]+" CHB",math.ceil(with_waste(net*self.a["block_per_m2"],self.a["chb_waste"])),"pcs")
        mortar=net*self.a["mortar_m3_per_m2"]; self.add("Masonry",r["ID"]+" mortar",mortar,"m³")
        plaster_area=net*(2 if yes(r,"Both-side plaster") else 1); self.add("Finishes",r["ID"]+" plaster area",plaster_area,"m²"); self.add("Finishes",r["ID"]+" plaster mortar",plaster_area*self.a["plaster_m3_per_m2"],"m³")
    def opening(self,r): self.add("Openings",r["ID"]+" "+str(r.get("Type","")),n(r,"Qty")*n(r,"Width m")*n(r,"Height m"),"m²")
    def total_net_masonry_area(self):
        """Wall face area after openings, used as the transparent painting basis."""
        return sum(max(0,n(r,"Qty")*n(r,"Length m")*n(r,"Height m")-n(r,"Opening area m2")) for r in self.p.items.get("Walls",[]))
    def painting(self,r):
        area=self.total_net_masonry_area()*n(r,"Painted sides")
        waste=n(r,"Waste %")
        coats=max(1,n(r,"Finish coats"))
        paint=with_waste(area*coats/max(.001,n(r,"Paint coverage m2/L",10)),waste)
        primer=with_waste(area/max(.001,n(r,"Primer coverage m2/L",10)),waste)
        self.add("Painting",r["ID"]+" masonry paint area",area,"m²")
        self.add("Painting",r["ID"]+" finish paint (%g coats)"%coats,paint,"L")
        self.add("Painting",r["ID"]+" primer",primer,"L")
    def tiles(self,r):
        area=n(r,"Area m2"); tile_area=(n(r,"Tile length mm")/1000)*(n(r,"Tile width mm")/1000)
        if tile_area <= 0: return
        pieces=math.ceil(with_waste(area/tile_area,n(r,"Waste %")))
        self.add("Tiles",r["ID"]+" "+str(r.get("Location",""))+" tile coverage",area,"m²")
        self.add("Tiles",r["ID"]+" %g×%g mm tiles"%(n(r,"Tile length mm"),n(r,"Tile width mm")),pieces,"pcs")
        self.add("Tiles",r["ID"]+" tile adhesive",with_waste(area*n(r,"Adhesive kg/m2"),n(r,"Waste %")),"kg")
        self.add("Tiles",r["ID"]+" grout",with_waste(area*n(r,"Grout kg/m2"),n(r,"Waste %")),"kg")
    def finish(self,r): self.add("Finishes",r["ID"]+" "+str(r.get("Type","")),with_waste(n(r,"Area m2")/max(.0001,n(r,"Coverage m2/unit",1)),n(r,"Waste %")),"units")

class ItemEditor(tk.Toplevel):
    def __init__(self, app, category, row=None):
        super().__init__(app); self.app=app; self.category=category; self.row=row; self.title(("Edit " if row else "Add ")+category); self.resizable(False,True); self.vars={}
        box=ttk.Frame(self,padding=12); box.grid(sticky="nsew")
        for i,(key,default) in enumerate(SCHEMAS[category]):
            ttk.Label(box,text=key+":").grid(row=i,column=0,sticky="e",padx=(0,8),pady=3)
            v=tk.StringVar(value=str(row.get(key,default) if row else default)); self.vars[key]=v
            ttk.Entry(box,textvariable=v,width=28).grid(row=i,column=1,sticky="ew",pady=3)
        ttk.Button(box,text="Save",command=self.save).grid(row=len(SCHEMAS[category]),column=0,columnspan=2,pady=(12,0))
        self.grab_set(); self.transient(app)
    def save(self):
        item={k:v.get().strip() for k,v in self.vars.items()}
        for key,default in SCHEMAS[self.category]:
            if isinstance(default,(int,float)) and f(item,key,-1)<0: messagebox.showerror("Invalid value",key+" cannot be negative.",parent=self); return
        if self.row: self.row.clear(); self.row.update(item)
        else: self.app.project.items[self.category].append(item)
        self.app.refresh(); self.destroy()

class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title(APP_TITLE); self.geometry("1280x760"); self.minsize(1000,650); self.project=self.sample(); self.current="Dashboard"; self.info_vars={}; self.make_style(); self.build(); self.refresh()
    def sample(self):
        p=Project()
        for cat in SCHEMAS:
            if cat in ("Footings","Columns","Tie Beams","Slabs","Suspended Slabs","Walls","Openings","Painting","Tiles","Finishes"): p.items[cat]=[{k:v for k,v in SCHEMAS[cat]}]
        return p
    def make_style(self):
        s=ttk.Style(self); s.theme_use("clam"); s.configure("Nav.TButton",padding=10,anchor="w"); s.configure("Title.TLabel",font=("Segoe UI",18,"bold")); s.configure("Card.TLabel",font=("Segoe UI",12,"bold")); s.configure("Treeview",rowheight=26)
    def build(self):
        outer=ttk.Frame(self); outer.pack(fill="both",expand=True)
        nav=ttk.Frame(outer,width=200,padding=10); nav.pack(side="left",fill="y"); nav.pack_propagate(False)
        ttk.Label(nav,text="BOM TAKEOFF",style="Title.TLabel").pack(anchor="w",pady=(4,18))
        for x in ["Dashboard","Project Information","Footings","Columns","Tie Beams","Elevated Beams","Slabs","Suspended Slabs","Walls","Openings","Painting","Tiles","Finishes","Results","Calculation Details"]:
            ttk.Button(nav,text=x,style="Nav.TButton",command=lambda z=x:self.show(z)).pack(fill="x",pady=2)
        ttk.Separator(nav).pack(fill="x",pady=12); ttk.Button(nav,text="Save project…",command=self.save_project).pack(fill="x",pady=2); ttk.Button(nav,text="Load project…",command=self.load_project).pack(fill="x",pady=2); ttk.Button(nav,text="Export Excel…",command=self.export_excel).pack(fill="x",pady=2)
        self.main=ttk.Frame(outer,padding=16); self.main.pack(side="left",fill="both",expand=True)
        self.disclaimer=tk.Label(self,text="PRELIMINARY ESTIMATING ONLY — Structural design, reinforcement sizing, loads, detailing and final NSCP 2015 code compliance require review and approval by a licensed structural engineer.",bg="#fff3cd",fg="#664d03",wraplength=950,justify="left",padx=10,pady=8); self.disclaimer.pack(side="bottom",fill="x")
    def show(self,page): self.current=page; self.refresh()
    def clear(self):
        for w in self.main.winfo_children(): w.destroy()
    def heading(self,text,color="#1f6fb2"):
        ttk.Label(self.main,text=text,style="Title.TLabel",foreground=color).pack(anchor="w",pady=(0,8))
    def refresh(self):
        self.clear()
        if self.current=="Dashboard": self.dashboard()
        elif self.current=="Project Information": self.project_info()
        elif self.current=="Calculation Details": self.details()
        elif self.current=="Results": self.results_page()
        else: self.table_page(self.current)
    def totals(self):
        lines=Calculator(self.project).calc(); d=defaultdict(float)
        for cat,desc,material,q,u in lines: d[(material,u)]+=q
        return lines,d

    def result_tree(self, parent, rows, height=16):
        """Reusable calculation-results table, with units kept alongside values."""
        tree=ttk.Treeview(parent,columns=("category","description","material","quantity","unit"),show="headings",height=height)
        for col, title, width in [("category","Category",125),("description","Element / calculation",250),
                                  ("material","Material",175),("quantity","Quantity",120),("unit","Unit",65)]:
            tree.heading(col,text=title); tree.column(col,width=width,anchor="w" if col in ("description","material") else "center")
        tree.pack(fill="both",expand=True,padx=4,pady=4)
        for cat, desc, material, qty, unit in rows:
            tree.insert("","end",values=(cat,desc,material,f"{qty:,.3f}",unit))
        return tree
    def dashboard(self):
        self.heading("Building Summary Dashboard","#198754"); lines,d=self.totals()
        def get(material,unit): return d.get((material,unit),0)
        metrics=[("Concrete",sum(q for c,_,_,q,u in lines if c=="Concrete"),"m³","#dbeafe"),("Reinforcing steel",sum(q for c,_,m,q,u in lines if c=="Reinforcement" and u=="kg"),"kg","#d1fae5"),("Formwork",sum(q for c,_,_,q,u in lines if c=="Formwork"),"m²","#dbeafe"),("CHB / blocks",sum(q for c,_,_,q,u in lines if c=="Masonry" and u=="pcs"),"pcs","#ffedd5"),("Net wall area",sum(q for c,desc,_,q,u in lines if "net wall area" in desc),"m²","#ffedd5"),("Floor area",sum(q for c,_,_,q,u in lines if c=="Flooring" and u=="m²"),"m²","#ffedd5")]
        cards=ttk.Frame(self.main); cards.pack(fill="x")
        for label,value,unit,bg in metrics:
            fr=tk.Frame(cards,bg=bg,padx=14,pady=12); fr.pack(side="left",padx=(0,10),pady=5,fill="both",expand=True); tk.Label(fr,text=label,bg=bg).pack(anchor="w"); tk.Label(fr,text=f"{value:,.2f} {unit}",font=("Segoe UI",14,"bold"),bg=bg).pack(anchor="w")
        ttk.Label(self.main,text="Major material summary",style="Card.TLabel").pack(anchor="w",pady=(22,4)); tree=ttk.Treeview(self.main,columns=("mat","qty","unit"),show="headings",height=15)
        for col,tx,w in [("mat","Material",350),("qty","Quantity",150),("unit","Unit",80)]: tree.heading(col,text=tx); tree.column(col,width=w)
        tree.pack(fill="both",expand=True)
        for (mat,u),q in sorted(d.items()): tree.insert("","end",values=(mat,f"{q:,.3f}",u))
        ttk.Button(self.main,text="Open detailed calculation results",command=lambda:self.show("Results")).pack(anchor="e",pady=(8,0))
    def project_info(self):
        self.heading("Project Information","#1f6fb2"); outer=ttk.Frame(self.main); outer.pack(anchor="nw",fill="x"); self.info_vars={}
        keys=list(self.project.info.items())+[("Assumption: "+k,v) for k,v in self.project.assumptions.items()]
        for i,(k,v) in enumerate(keys):
            ttk.Label(outer,text=k+":").grid(row=i//2,column=(i%2)*2,sticky="e",padx=6,pady=6); var=tk.StringVar(value=str(v)); self.info_vars[k]=var; ent=ttk.Entry(outer,textvariable=var,width=34); ent.grid(row=i//2,column=(i%2)*2+1,sticky="w",padx=6,pady=6)
        ttk.Button(self.main,text="Apply project settings",command=self.apply_info).pack(anchor="w",pady=16)
        ttk.Label(self.main,text="Waste factors are percentages. Concrete mix references are editable preliminary assumptions, not a structural design specification.",wraplength=800).pack(anchor="w")
    def apply_info(self):
        for k,v in self.info_vars.items():
            if k.startswith("Assumption: "):
                key=k[12:]; value=f({key:v.get()},key); self.project.assumptions[key]=value
            else:self.project.info[k]=v.get()
        self.refresh()
    def table_page(self,cat):
        color="#1f6fb2" if cat in ("Footings","Columns","Tie Beams","Elevated Beams","Slabs","Suspended Slabs") else "#e67e22"
        self.heading(cat,color); ttk.Label(self.main,text="Add multiple member types/locations. Select a row to edit, duplicate or delete it.").pack(anchor="w",pady=(0,8))
        cols=[x[0] for x in SCHEMAS[cat]]; tree=ttk.Treeview(self.main,columns=cols,show="headings",height=20)
        for c in cols: tree.heading(c,text=c); tree.column(c,width=max(90,min(150,len(c)*9)),anchor="center")
        tree.pack(fill="both",expand=True); self.tree=tree
        for i,row in enumerate(self.project.items[cat]): tree.insert("","end",iid=str(i),values=[row.get(c,"") for c in cols])
        bar=ttk.Frame(self.main); bar.pack(fill="x",pady=10)
        ttk.Button(bar,text="+ Add",command=lambda:ItemEditor(self,cat)).pack(side="left",padx=(0,6)); ttk.Button(bar,text="Edit selected",command=lambda:self.edit_selected(cat)).pack(side="left",padx=6); ttk.Button(bar,text="Duplicate",command=lambda:self.duplicate(cat)).pack(side="left",padx=6); ttk.Button(bar,text="Delete",command=lambda:self.delete(cat)).pack(side="left",padx=6); ttk.Button(bar,text="View results",command=lambda:self.show("Results")).pack(side="right")
        diagram_box=ttk.LabelFrame(self.main,text="Graphical member guide — select a table row to update labels",padding=5)
        diagram_box.pack(fill="x",pady=(2,0)); self.diagram_canvas=tk.Canvas(diagram_box,height=145,bg="#f8fafc",highlightthickness=0); self.diagram_canvas.pack(side="left",fill="x",expand=True)
        ttk.Label(diagram_box,text=self.diagram_text(cat),foreground="#555",justify="left",wraplength=290).pack(side="right",padx=10)
        tree.bind("<<TreeviewSelect>>",lambda e:self.draw_diagram(cat))
        self.draw_diagram(cat)

    def draw_diagram(self, cat):
        """Simple not-to-scale construction sketches with the selected member's dimensions."""
        cv=getattr(self,"diagram_canvas",None)
        if not cv:return
        cv.delete("all"); row={}
        sel=self.tree.selection() if hasattr(self,"tree") else ()
        if sel: row=self.project.items[cat][int(sel[0])]
        elif self.project.items.get(cat): row=self.project.items[cat][0]
        cv.create_text(12,12,anchor="w",text="Selected: "+str(row.get("ID","Add a member to view diagram")),font=("Segoe UI",10,"bold"),fill="#1f2937")
        ink="#1f6fb2"; dim="#e67e22"; x0,y0,x1,y1=115,40,315,115
        def label(x,y,text): cv.create_text(x,y,text=text,fill=dim,font=("Segoe UI",9,"bold"))
        def arrow(xa,ya,xb,yb,text):
            cv.create_line(xa,ya,xb,yb,fill=dim,arrow="both"); label((xa+xb)/2,(ya+yb)/2-10,text)
        if cat=="Footings":
            cv.create_rectangle(x0,y0,x1,y1,outline=ink,width=3,fill="#dbeafe"); cv.create_line(x0,y0+25,x1,y0+25,fill=ink,dash=(4,3)); arrow(x0,130,x1,130,"L = %s m"%row.get("Length m","—")); arrow(85,y0,85,y1,"W = %s m"%row.get("Width m","—")); label(380,62,"t = %s m"%row.get("Thick m","—")); label(380,87,"bottom bars X/Y")
        elif cat=="Columns":
            cv.create_rectangle(175,35,255,120,outline=ink,width=3,fill="#dbeafe"); arrow(270,35,270,120,"H = %s m"%row.get("Height/storey m","—")); arrow(175,130,255,130,"W = %s m"%row.get("Width m","—")); label(360,62,"D = %s m"%row.get("Depth m","—")); label(360,87,"main bars + ties")
        elif cat in ("Tie Beams","Elevated Beams"):
            cv.create_rectangle(85,62,390,110,outline=ink,width=3,fill="#dbeafe");
            for x in range(105,385,25): cv.create_line(x,65,x,107,fill="#64748b")
            arrow(85,130,390,130,"L = %s m"%row.get("Length m","—")); arrow(410,62,410,110,"D = %s m"%row.get("Depth m","—")); label(470,72,"W = %s m"%row.get("Width m","—")); label(470,98,"top / bottom bars")
        elif cat in ("Slabs","Suspended Slabs"):
            cv.create_polygon(115,50,290,30,390,85,210,108,outline=ink,fill="#dbeafe",width=3)
            for x in range(150,350,32): cv.create_line(x,43,x+85,48,fill="#64748b")
            arrow(115,125,390,125,"L = %s m"%row.get("Length m","—")); label(410,55,"W = %s m"%row.get("Width m","—")); label(410,78,"t = %s m"%row.get("Thickness m","—")); label(410,101,"top + bottom bars X / Y" if cat=="Suspended Slabs" else "reinforcement X / Y")
        else:
            cv.create_rectangle(125,35,355,115,outline=ink,width=3,fill="#ffedd5"); arrow(125,130,355,130,"L / W = %s m"%row.get("Length m",row.get("Width m","—"))); arrow(105,35,105,115,"H = %s m"%row.get("Height m","—")); label(390,60,"Qty = %s"%row.get("Qty","—")); label(390,85,"area-based takeoff")
    def diagram_text(self,c):
        return {"Footings":"▭ Footing schematic: L × W × thickness; bars each way at entered spacing.","Columns":"▯ Column schematic: width × depth × storey height; vertical bars with ties.","Tie Beams":"━ Beam schematic: length × width × depth; top/bottom bars and stirrups.","Elevated Beams":"━ Elevated beam: use clear member length for preliminary takeoff.","Slabs":"▱ Slab-on-grade: length × width × thickness; bars in X and Y directions.","Suspended Slabs":"▱ Suspended slab: top and bottom reinforcement mats, soffit formwork and slab-edge formwork.","Walls":"▯ Wall: gross area less entered openings; plaster may be one or both sides."}.get(c,"▯ Opening/finish dimensions are taken directly from table entries.")
    def selected(self):
        sel=self.tree.selection()
        if not sel: messagebox.showinfo("Select a row","Please select a row first."); return None
        return int(sel[0])
    def edit_selected(self,cat):
        i=self.selected()
        if i is not None: ItemEditor(self,cat,self.project.items[cat][i])
    def duplicate(self,cat):
        i=self.selected()
        if i is not None:self.project.items[cat].append(self.project.items[cat][i].copy()); self.refresh()
    def delete(self,cat):
        i=self.selected()
        if i is not None and messagebox.askyesno("Delete","Delete selected item?"): self.project.items[cat].pop(i); self.refresh()
    def results_page(self):
        """Results are intentionally separated by construction element before the grand list."""
        self.heading("Calculation Results","#198754")
        ttk.Label(self.main,text="Quantities update from the current input tables. Review assumptions and calculations before relying on any estimate.").pack(anchor="w",pady=(0,7))
        lines,_ = self.totals()
        book=ttk.Notebook(self.main); book.pack(fill="both",expand=True)
        overall=ttk.Frame(book); book.add(overall,text="Overall materials")
        grouped=defaultdict(float)
        for cat,desc,mat,q,u in lines:
            # Aggregated steel already represents the material total; raw rebar lengths remain a separate m line.
            grouped[(cat,mat,u)]+=q
        summary=[(cat,"All applicable elements",mat,q,u) for (cat,mat,u),q in sorted(grouped.items())]
        self.result_tree(overall,summary)
        structural=["Footings","Columns","Tie Beams","Elevated Beams","Slabs","Suspended Slabs"]
        for cat in structural:
            tab=ttk.Frame(book); book.add(tab,text=cat)
            ids=[str(r.get("ID","")) for r in self.project.items.get(cat,[])]
            rows=[]
            for line in lines:
                line_cat,desc,mat,q,u=line
                if line_cat=="Reinforcement" and desc in ("Aggregated bars","Stock bars"): continue
                if any(desc.startswith(member_id) for member_id in ids if member_id): rows.append(line)
            if not rows:
                ttk.Label(tab,text="No calculated items yet. Add a member in the "+cat+" section.").pack(padx=18,pady=18,anchor="w")
            else: self.result_tree(tab,rows)
        arch=ttk.Frame(book); book.add(arch,text="Architectural")
        arch_rows=[line for line in lines if line[0] in ("Masonry","Openings","Painting","Tiles","Finishes","Flooring")]
        self.result_tree(arch,arch_rows) if arch_rows else ttk.Label(arch,text="No architectural quantities yet.").pack(padx=18,pady=18,anchor="w")
    def details(self):
        self.heading("Calculation Details & Assumptions","#198754"); text=tk.Text(self.main,wrap="word",font=("Consolas",10)); text.pack(fill="both",expand=True)
        content="""FORMULAS (all dimensions in metres after mm ÷ 1000 conversion)

Concrete: member geometric volume × (1 + concrete waste factor).
Footing excavation: (length + 2 working space) × (width + 2 working space) × excavation depth × quantity.
Columns: width × depth × storey height × quantity × storeys.
Beams: length × width × depth × quantity. Slabs: length × width × thickness × floors.
Formwork: exposed side/soffit geometric areas × (1 + formwork waste).
Rebar: count × run length, then weight = length × listed unit weight. Rebar is aggregated by diameter, including the editable rebar waste. Stock bars are ceiling(total length / configured stock length).
Walls: net area = quantity × length × height − entered opening area. CHB = net area × configurable blocks/m² × waste. Plaster is net wall area × 1 or 2 sides.

LIMITATION
This program is solely a preliminary quantity-takeoff/estimating tool. It does not design members, choose reinforcement, evaluate loads, detail bars, or certify compliance. NSCP 2015 terms and editable reference fields are provided only for context. A licensed structural engineer must review and approve all final design, drawings, quantities, and code compliance.

CURRENT ASSUMPTIONS
"""+"\n".join(f"• {k}: {v}" for k,v in self.project.assumptions.items())
        text.insert("1.0",content); text.config(state="disabled")
    def save_project(self):
        p=filedialog.asksaveasfilename(defaultextension=".json",filetypes=[("JSON project","*.json")])
        if p:
            with open(p,"w",encoding="utf-8") as fh: json.dump(asdict(self.project),fh,indent=2)
    def load_project(self):
        p=filedialog.askopenfilename(filetypes=[("JSON project","*.json")])
        if p:
            try:
                with open(p,encoding="utf-8") as fh: data=json.load(fh)
                self.project=Project(**data)
                for category in SCHEMAS: self.project.items.setdefault(category,[])
                self.refresh()
            except Exception as e: messagebox.showerror("Cannot load project",str(e))
    def export_excel(self):
        if not EXCEL_AVAILABLE: messagebox.showerror("Excel export unavailable","Install openpyxl first:\npy -m pip install openpyxl"); return
        p=filedialog.asksaveasfilename(defaultextension=".xlsx",filetypes=[("Excel workbook","*.xlsx")])
        if not p:return
        try: self.make_workbook(p); messagebox.showinfo("Export complete","Workbook saved:\n"+p)
        except Exception as e: messagebox.showerror("Export failed",str(e))
    def make_workbook(self,path):
        wb=Workbook(); cover=wb.active; cover.title="Cover"
        blue="1F6FB2"; orange="E67E22"; green="198754"; header=PatternFill("solid",fgColor=blue)
        cover.append([APP_TITLE,VERSION]); cover["A1"].font=Font(bold=True,size=16,color="FFFFFF"); cover["A1"].fill=header; cover.merge_cells("A1:D1")
        for k,v in self.project.info.items(): cover.append([k,v])
        cover.append([]); cover.append(["DISCLAIMER","Preliminary estimating only. Final design and NSCP 2015 compliance require licensed structural engineer review."])
        for cat, rows in self.project.items.items():
            ws=wb.create_sheet(cat[:31]); cols=[x[0] for x in SCHEMAS[cat]]; ws.append(cols)
            for cell in ws[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=blue if cat in ("Footings","Columns","Tie Beams","Elevated Beams","Slabs","Suspended Slabs") else orange)
            for row in rows: ws.append([row.get(c,"") for c in cols]); ws.freeze_panes="A2"; self.autosize(ws)
        lines,_=self.totals(); ws=wb.create_sheet("Overall Material Summary"); ws.append(["Category","Description","Material","Quantity","Unit"])
        for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor=green)
        for line in lines:ws.append(line)
        ws.freeze_panes="A2"; self.autosize(ws)
        for name,cats in [("Concrete Summary",["Concrete","Concrete materials"]),("Reinforcement Summary",["Reinforcement"]),("Formwork Summary",["Formwork"]),("Masonry Wall Summary",["Masonry"]),("Finishes Summary",["Painting","Tiles","Finishes","Flooring"])]:
            sh=wb.create_sheet(name); sh.append(["Description","Material","Quantity","Unit"])
            for cell in sh[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor=green)
            for c,d,m,q,u in lines:
                if c in cats: sh.append([d,m,q,u])
            sh.freeze_panes="A2"; self.autosize(sh)
        ass=wb.create_sheet("Calculation Assumptions"); ass.append(["Assumption","Value"])
        for k,v in self.project.assumptions.items(): ass.append([k,v])
        ass.append([]);ass.append(["Professional limitation","Preliminary takeoff only; no independent NSCP compliance certification."]); self.autosize(ass)
        wb.save(path)
    @staticmethod
    def autosize(ws):
        for col in ws.columns:
            letter=col[0].column_letter; ws.column_dimensions[letter].width=min(45,max(12,max(len(str(c.value or "")) for c in col)+2))

if __name__ == "__main__":
    App().mainloop()
